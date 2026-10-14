"""
Document Ingestion — Загрузка и парсинг документов
Поддержка PDF, TXT, MD файлов
"""
import asyncio
import logging
from pathlib import Path
from typing import List, Dict, Optional
import hashlib

logger = logging.getLogger(__name__)


class Document:
    """Представление документа"""
    
    def __init__(self, content: str, metadata: Dict):
        self.content = content
        self.metadata = metadata
        self.id = self._generate_id()
    
    def _generate_id(self) -> str:
        """Генерация уникального ID документа"""
        content_hash = hashlib.md5(self.content.encode()).hexdigest()[:16]
        source = self.metadata.get('source', 'unknown')
        return f"{source}_{content_hash}"
    
    def __repr__(self):
        return f"Document(id={self.id}, source={self.metadata.get('source', 'unknown')}, length={len(self.content)})"


class DocumentIngestion:
    """
    Загрузка и парсинг документов
    Поддержка форматов: PDF, TXT, MD
    """
    
    def __init__(self, config: Dict):
        self.config = config
        self.documents_path = Path(config.get('documents_path', './documents'))
        self.supported_formats = ['.pdf', '.txt', '.md', '.docx']
        
        logger.info(f"DocumentIngestion инициализирован (путь: {self.documents_path})")
    
    async def ingest_directory(self, directory: Optional[Path] = None) -> List[Document]:
        """
        Загрузка всех документов из директории
        
        Args:
            directory: Путь к директории (по умолчанию self.documents_path)
        
        Returns:
            Список загруженных документов
        """
        if directory is None:
            directory = self.documents_path
        
        if not directory.exists():
            logger.warning(f"Директория не существует: {directory}")
            return []
        
        documents = []
        
        # Рекурсивный обход директории
        for file_path in directory.rglob('*'):
            if file_path.is_file() and file_path.suffix.lower() in self.supported_formats:
                try:
                    doc = await self.ingest_file(file_path)
                    if doc:
                        documents.append(doc)
                        logger.info(f"✅ Загружен: {file_path.name}")
                except Exception as e:
                    logger.error(f"❌ Ошибка загрузки {file_path}: {e}")
        
        logger.info(f"📚 Загружено документов: {len(documents)}")
        return documents
    
    async def ingest_file(self, file_path: Path) -> Optional[Document]:
        """
        Загрузка одного документа
        
        Args:
            file_path: Путь к файлу
        
        Returns:
            Document или None
        """
        suffix = file_path.suffix.lower()
        
        if suffix == '.pdf':
            return await self._ingest_pdf(file_path)
        elif suffix == '.txt':
            return await self._ingest_txt(file_path)
        elif suffix == '.md':
            return await self._ingest_markdown(file_path)
        elif suffix == '.docx':
            return await self._ingest_docx(file_path)
        else:
            logger.warning(f"Неподдерживаемый формат: {suffix}")
            return None
    
    async def _ingest_pdf(self, file_path: Path) -> Optional[Document]:
        """Загрузка PDF файла"""
        try:
            # Попытка использовать PyPDF2
            try:
                import PyPDF2
                
                def extract_pdf():
                    with open(file_path, 'rb') as f:
                        reader = PyPDF2.PdfReader(f)
                        text = ""
                        for page in reader.pages:
                            text += page.extract_text() + "\n"
                        return text
                
                content = await asyncio.to_thread(extract_pdf)
                
                metadata = {
                    'source': str(file_path),
                    'filename': file_path.name,
                    'format': 'pdf',
                    'pages': len(PyPDF2.PdfReader(open(file_path, 'rb')).pages),
                }
                
                return Document(content, metadata)
            
            except ImportError:
                logger.warning("PyPDF2 не установлен. Установите: pip install PyPDF2")
                return None
        
        except Exception as e:
            logger.error(f"Ошибка загрузки PDF {file_path}: {e}")
            return None
    
    async def _ingest_txt(self, file_path: Path) -> Optional[Document]:
        """Загрузка TXT файла"""
        try:
            content = await asyncio.to_thread(file_path.read_text, encoding='utf-8')
            
            metadata = {
                'source': str(file_path),
                'filename': file_path.name,
                'format': 'txt',
            }
            
            return Document(content, metadata)
        
        except Exception as e:
            logger.error(f"Ошибка загрузки TXT {file_path}: {e}")
            return None
    
    async def _ingest_markdown(self, file_path: Path) -> Optional[Document]:
        """Загрузка Markdown файла"""
        try:
            content = await asyncio.to_thread(file_path.read_text, encoding='utf-8')
            
            metadata = {
                'source': str(file_path),
                'filename': file_path.name,
                'format': 'md',
            }
            
            return Document(content, metadata)
        
        except Exception as e:
            logger.error(f"Ошибка загрузки Markdown {file_path}: {e}")
            return None
    
    async def _ingest_docx(self, file_path: Path) -> Optional[Document]:
        """Загрузка DOCX файла"""
        try:
            # Попытка использовать python-docx
            try:
                import docx
                
                def extract_docx():
                    doc = docx.Document(file_path)
                    text = ""
                    for paragraph in doc.paragraphs:
                        text += paragraph.text + "\n"
                    return text
                
                content = await asyncio.to_thread(extract_docx)
                
                metadata = {
                    'source': str(file_path),
                    'filename': file_path.name,
                    'format': 'docx',
                }
                
                return Document(content, metadata)
            
            except ImportError:
                logger.warning("python-docx не установлен. Установите: pip install python-docx")
                return None
        
        except Exception as e:
            logger.error(f"Ошибка загрузки DOCX {file_path}: {e}")
            return None
    
    async def ingest_text(self, text: str, source: str = "manual") -> Document:
        """
        Загрузка текста напрямую
        
        Args:
            text: Текст документа
            source: Источник (для метаданных)
        
        Returns:
            Document
        """
        metadata = {
            'source': source,
            'format': 'text',
        }
        
        return Document(text, metadata)
